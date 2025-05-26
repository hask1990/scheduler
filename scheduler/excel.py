"""Excel assignment scheduler.

This module provides a Python translation of an Office Script that assigns teams
to time slots while skipping blackout weeks defined by consecutive holidays.
The script uses :mod:`openpyxl` to manipulate Excel workbooks.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta
from typing import Iterable, List, Set

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import range_boundaries


@dataclass
class Config:
    """Configuration for :func:`assign`."""

    teams_range_address: str = "A2:A22"
    holidays_sheet_name: str = "Holidays"
    assignments_sheet_name: str = "Assignments"
    start_hour: int = 9
    slot_hours: int = 1
    workday_end_hour: int = 17
    schedule_sheet_name: str = "Schedule"


def _get_iso_week(date: datetime) -> int:
    dt = datetime(date.year, date.month, date.day)
    dt += timedelta(days=4 - (dt.isoweekday() or 7))
    year_start = datetime(dt.year, 1, 1)
    return ((dt - year_start).days + 1 + 6) // 7


def _load_teams(sheet: Worksheet, cell_range: str) -> List[str]:
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    teams = []
    for row in sheet.iter_rows(min_row=min_row, max_row=max_row,
                               min_col=min_col, max_col=max_col):
        for cell in row:
            if cell.value:
                teams.append(str(cell.value))
    return teams


def _load_holidays(sheet: Worksheet) -> List[datetime]:
    dates = []
    for cell in sheet.iter_cols(min_col=1, max_col=1, min_row=2, values_only=True)[0]:
        if cell:
            if isinstance(cell, datetime):
                dates.append(cell)
            else:
                dates.append(datetime.fromisoformat(str(cell)))
    return sorted(dates)


def _compute_blackout_weeks(dates: Iterable[datetime]) -> Set[int]:
    dates = sorted(dates)
    blackout_weeks: Set[int] = set()
    i = 0
    while i < len(dates):
        run = [dates[i]]
        j = i + 1
        while j < len(dates) and (dates[j] - dates[j - 1]).days == 1:
            run.append(dates[j])
            j += 1
        if len(run) >= 3:
            for d in run:
                blackout_weeks.add(_get_iso_week(d))
        i = j
    return blackout_weeks


def assign(workbook_path: str, *, config: Config = Config()) -> None:
    """Fill the assignments sheet with generated slots."""

    wb = load_workbook(workbook_path)
    schedule_sheet = wb[config.schedule_sheet_name]
    holidays_sheet = wb[config.holidays_sheet_name]
    assignments_sheet = wb[config.assignments_sheet_name]

    teams = _load_teams(schedule_sheet, config.teams_range_address)
    holidays = _load_holidays(holidays_sheet)
    blackout = _compute_blackout_weeks(holidays)

    available_dates: List[datetime] = []
    cursor = datetime.now()
    while len(available_dates) < len(teams):
        if cursor.weekday() < 5 and _get_iso_week(cursor) not in blackout:
            available_dates.append(cursor.replace(hour=0, minute=0, second=0, microsecond=0))
        cursor += timedelta(days=1)

    hour = config.start_hour
    day_idx = 0
    rows: List[List] = []
    for team in teams:
        date = available_dates[day_idx]
        rows.append([team, date.date(), hour])
        hour += config.slot_hours
        if hour > config.workday_end_hour:
            hour = config.start_hour
            day_idx += 1

    assignments_sheet.delete_rows(2, assignments_sheet.max_row)
    for row in rows:
        assignments_sheet.append(row)

    wb.save(workbook_path)


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Generate team assignments in an Excel workbook")
    parser.add_argument("workbook", help="Path to the workbook")
    args = parser.parse_args()

    assign(args.workbook)
